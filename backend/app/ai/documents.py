"""Reading a document about a student and finding what they need help with.

The gradebook engine knows the numbers. A teacher's note, a report card comment
or an essay holds evidence it cannot see: "loses track partway through
multi-step problems", "stops writing once she runs out of ideas". This module
turns that prose into specific, checkable findings.

A 4B model reading prose will invent things, so every finding has to survive
three deterministic checks before anyone sees it:

1. **It must quote the document.** Each finding carries a verbatim quotation, and
   the quotation must actually occur in the extracted text. A need the model made
   up has nothing to quote, and is dropped — the same principle as refusing an
   agent's invented SKU.
2. **It may not diagnose.** "Struggles to start written tasks" is an observation a
   support office can act on. "Has dyslexia" is a clinical judgement that belongs
   to qualified people; a finding that asserts one is withheld with a notice.
3. **It is checked against the gradebook independently.** The model never sees the
   student's grades, so it cannot parrot them back as findings. Corroboration is
   computed afterwards: where the document and the gradebook agree the finding is
   stronger, and where they disagree that is worth a person's attention too.

Context is handled explicitly. Ollama runs this model at a 4,096-token window
unless told otherwise and cuts overflow silently, so documents are split into
chunks well inside the configured window, and every reply is checked for signs of
truncation.
"""
from __future__ import annotations

import hashlib
import io
import json
import re
import time
from dataclasses import dataclass
from datetime import datetime

from sqlalchemy.orm import Session

from ..analytics import StudentSignal, build_signals
from ..config import get_settings
from ..models import StudentDocument
from .ollama import OllamaError, chat
from .tools import PLAN_KINDS

settings = get_settings()

MAX_BYTES = 10 * 1024 * 1024
CHUNK_CHARS = 9000           # ~2.5k tokens of prose; the prompt stays far inside the window
MAX_CHUNKS = 12              # beyond ~100k characters, ask for the relevant section instead
MAX_FINDINGS_PER_CHUNK = 8
KINDS = ["report-card", "teacher-note", "essay", "assessment", "counselor-note", "other"]

EXTENSIONS = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".txt": "text/plain",
    ".md": "text/markdown",
}


class DocumentError(ValueError):
    """The upload cannot be read. The message is shown to the person uploading it."""


# ---------------------------------------------------------------------------
# extraction
# ---------------------------------------------------------------------------
def extension_of(filename: str) -> str:
    m = re.search(r"(\.[A-Za-z0-9]+)$", filename or "")
    return m.group(1).lower() if m else ""


def extract_text(filename: str, data: bytes) -> tuple[str, int | None]:
    """Return (text, page count). Raises DocumentError with a usable explanation."""
    ext = extension_of(filename)
    if ext not in EXTENSIONS:
        raise DocumentError(
            f"{filename or 'That file'} is not a supported type. Upload a PDF, a Word document "
            "(.docx), or a text file.")
    if len(data) > MAX_BYTES:
        raise DocumentError(f"That file is {len(data) / 1_048_576:.1f} MB; the limit is 10 MB.")
    if not data:
        raise DocumentError("That file is empty.")

    # Check the contents match the name, rather than trusting an extension.
    if ext == ".pdf" and not data.lstrip()[:5].startswith(b"%PDF"):
        raise DocumentError(f"{filename} is named .pdf but is not a PDF.")
    if ext == ".docx" and not data[:2] == b"PK":
        raise DocumentError(f"{filename} is named .docx but is not a Word document.")

    pages: int | None = None
    try:
        if ext == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(data))
            if reader.is_encrypted:
                raise DocumentError(f"{filename} is password-protected. Remove the password and try again.")
            pages = len(reader.pages)
            text = "\n\n".join((page.extract_text() or "") for page in reader.pages)
        elif ext == ".docx":
            import docx

            document = docx.Document(io.BytesIO(data))
            parts = [p.text for p in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    parts.append(" | ".join(cell.text.strip() for cell in row.cells))
            text = "\n".join(parts)
        else:
            text = data.decode("utf-8-sig", errors="replace")
    except DocumentError:
        raise
    except Exception as e:
        raise DocumentError(f"{filename} could not be read: {type(e).__name__}.") from e

    text = re.sub(r"[ \t ]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if len(text) < 40:
        hint = (" It looks like a scan or a photo — this setup has no vision model, so only "
                "documents containing real text can be read.") if ext == ".pdf" else ""
        raise DocumentError(f"{filename} contains no readable text.{hint}")
    return text, pages


def chunk_text(text: str, size: int = CHUNK_CHARS) -> list[str]:
    """Split on paragraph boundaries; hard-split only a paragraph longer than a chunk."""
    paragraphs = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    chunks: list[str] = []
    current = ""
    for para in paragraphs:
        while len(para) > size:
            cut = para.rfind(". ", 0, size)
            cut = cut + 1 if cut > size // 2 else size
            if current:
                chunks.append(current)
                current = ""
            chunks.append(para[:cut].strip())
            para = para[cut:].strip()
        if len(current) + len(para) + 2 > size and current:
            chunks.append(current)
            current = para
        else:
            current = f"{current}\n\n{para}" if current else para
    if current:
        chunks.append(current)
    return chunks


def sha256_of(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


# ---------------------------------------------------------------------------
# the three checks
# ---------------------------------------------------------------------------
_QUOTE_MAP = str.maketrans({"‘": "'", "’": "'", "“": '"', "”": '"',
                            "–": "-", "—": "-", " ": " "})


def normalise(s: str) -> str:
    s = (s or "").translate(_QUOTE_MAP).lower()
    return re.sub(r"\s+", " ", s).strip()


def quote_in_document(quote: str, normalised_text: str) -> bool:
    """True when the quotation really occurs in the document.

    Case, curly quotes and line breaks are forgiven — PDF extraction mangles all
    three. An ellipsis is allowed to join two genuine fragments in order. Words are
    not forgiven: a paraphrase is not a quotation.
    """
    q = normalise(quote).strip(" \"'.,;:")
    if len(q) < 12:
        return False                      # "math" or "homework" proves nothing
    parts = [p.strip(" \"'.,;:") for p in re.split(r"\.\.\.|…", q)]
    parts = [p for p in parts if p]
    if not parts or any(len(p) < 6 for p in parts):
        return False
    position = 0
    for part in parts:
        found = normalised_text.find(part, position)
        if found < 0:
            return False
        position = found + len(part)
    return True


_DIAGNOSTIC = re.compile(
    r"\b(adhd|add\b|dyslexi\w*|dyscalculi\w*|dysgraphi\w*|autis\w*|asperger\w*|"
    r"bipolar|ocd|depress(ion|ed|ive)|anxiety disorder|learning disabilit\w*|"
    r"disorder|diagnos\w*|on the spectrum|neurodivergen\w*)\b", re.I)


def asserts_a_diagnosis(finding: dict) -> str | None:
    """The term that makes this a clinical claim, or None if it describes behaviour."""
    for field in ("area", "explanation"):
        m = _DIAGNOSTIC.search(str(finding.get(field) or ""))
        if m:
            return m.group(0)
    return None


# Everyday words a teacher writes, mapped to the classes they mean. A department is
# only used where every class in it fits the word: "math" is both maths classes, but
# "writing" is one class, not all of Humanities — mapping it to the department put
# Model UN's grade next to a finding about short stories.
SUBJECT_DEPARTMENTS = {
    "math": "Mathematics", "maths": "Mathematics", "algebra": "Mathematics",
    "geometry": "Mathematics", "equation": "Mathematics", "equations": "Mathematics",
    "science": "Science", "lab": "Science",
}
SUBJECT_TITLES = {
    "writing": "writing", "essay": "writing", "stories": "writing", "story": "writing",
    "spanish": "spanish", "debate": "model un",
    "photograph": "photography", "photographs": "photography", "photography": "photography",
    "ceramics": "ceramics", "clay": "ceramics", "band": "jazz band", "music": "jazz band",
    "robotics": "robotics", "cooking": "culinary", "kitchen": "culinary",
    "coding": "game design", "gaming": "esports", "climbing": "climbing", "forensic": "forensic",
}
ATTENDANCE_WORDS = ("absent", "absence", "attendance", "missed school", "tardy", "late to class", "not in class")
SUBMISSION_WORDS = ("homework", "missing", "turn in", "turned in", "hand in", "handed in",
                    "submit", "incomplete", "late work", "not completed", "unfinished", "assignments")


def _matched_courses(hay: str, sig: StudentSignal) -> list:
    named = [c for c in sig.courses
             if normalise(c.course_code) in hay or normalise(c.course_title) in hay]
    by_word = []
    for word, dept in SUBJECT_DEPARTMENTS.items():
        if re.search(rf"\b{re.escape(word)}\b", hay):
            by_word += [c for c in sig.courses if c.dept == dept]
    for word, title in SUBJECT_TITLES.items():
        if re.search(rf"\b{re.escape(word)}\b", hay):
            by_word += [c for c in sig.courses if title in c.course_title.lower()]
    out, seen = [], set()
    for c in named + by_word:
        if c.course_code not in seen:
            seen.add(c.course_code)
            out.append(c)
    return out


def corroborate(finding: dict, sig: StudentSignal | None) -> dict:
    """What the gradebook says about the same thing, computed without the model."""
    if sig is None:
        return {"verdict": "no-record", "evidence": []}
    hay = normalise(" ".join(str(finding.get(k) or "") for k in ("area", "quote", "explanation")))
    evidence: list[dict] = []
    courses = _matched_courses(hay, sig)

    for course in sig.courses:
        for strand in course.skills:
            if normalise(strand.skill) in hay:
                evidence.append({"source": f"{strand.skill} ({course.course_code})",
                                 "detail": f"{strand.pct:.0f}%", "pct": strand.pct, "kind": "strand"})
    for course in courses:
        evidence.append({"source": f"{course.course_title} ({course.course_code})",
                         "detail": f"{course.pct:.0f}%", "pct": course.pct, "kind": "course"})

    if any(w in hay for w in ATTENDANCE_WORDS) and sig.days_counted:
        evidence.append({"source": "attendance", "detail": f"absent {sig.absences} of {sig.days_counted} days",
                         "rate": sig.absence_rate, "kind": "attendance",
                         "concern": sig.absence_rate >= 0.10})
    if any(w in hay for w in SUBMISSION_WORDS):
        # About specific classes when the finding names them; the whole timetable otherwise.
        scope = courses or sig.courses
        missing = sum(c.missing for c in scope)
        graded = sum(c.graded_items for c in scope) or 1
        where = ", ".join(c.course_code for c in courses) if courses else "all classes"
        evidence.append({"source": f"work submitted ({where})",
                         "detail": f"{missing} of {graded} graded pieces missing",
                         "rate": round(missing / graded, 3), "kind": "missing",
                         "concern": missing / graded >= 0.10})

    if not evidence:
        return {"verdict": "no-evidence", "evidence": []}

    # Strands are the most specific evidence; judge on them when a strand matched.
    strand_pcts = [e["pct"] for e in evidence if e["kind"] == "strand"]
    pcts = strand_pcts or [e["pct"] for e in evidence if "pct" in e]
    concerns = [e["concern"] for e in evidence if "concern" in e]
    if finding.get("type") == "strength":
        agrees = any(p >= 85 for p in pcts)
        disagrees = bool(pcts) and all(p < settings.support_threshold for p in pcts)
    else:
        agrees = any(p < settings.support_threshold for p in pcts) or any(concerns)
        disagrees = bool(pcts) and all(p >= 85 for p in pcts) and not any(concerns)
    verdict = "agrees" if agrees and not disagrees else "disagrees" if disagrees else "mixed"
    return {"verdict": verdict, "evidence": evidence}


SEVERITY_ORDER = {"low": 0, "medium": 1, "high": 2}


def calibrate_severity(finding: dict) -> tuple[str, str | None]:
    """Raise a need's severity when the gradebook shows it is worse than the model said.

    Observed: the model rated "multi-step word problems" low for a student at 13% on
    that strand. Severity is only ever raised, never lowered — a teacher who writes
    about a serious concern may know something the grades do not yet show.
    """
    given = finding.get("severity") if finding.get("severity") in SEVERITY_ORDER else "medium"
    if finding.get("type") != "need":
        return given, None
    floor, why = "low", None
    for e in finding.get("gradebook", {}).get("evidence", []):
        if e.get("kind") in ("strand", "course") and e["pct"] < 50:
            floor, why = "high", f"{e['source']} is at {e['detail']}"
            break
        if e.get("kind") in ("attendance", "missing") and e.get("rate", 0) >= 0.25:
            floor, why = "high", f"{e['source']}: {e['detail']}"
            break
        if e.get("kind") in ("strand", "course") and e["pct"] < settings.support_threshold \
                and SEVERITY_ORDER[floor] < 1:
            floor, why = "medium", f"{e['source']} is at {e['detail']}"
        if e.get("kind") in ("attendance", "missing") and e.get("rate", 0) >= 0.10 \
                and SEVERITY_ORDER[floor] < 1:
            floor, why = "medium", f"{e['source']}: {e['detail']}"
    if SEVERITY_ORDER[floor] > SEVERITY_ORDER[given]:
        return floor, f"Raised from {given}: {why}."
    return given, None


# ---------------------------------------------------------------------------
# the model call
# ---------------------------------------------------------------------------
SUPPORT_OPTIONS = PLAN_KINDS + ["none"]

SCHEMA = {
    "type": "object",
    "properties": {
        "findings": {
            "type": "array",
            "items": {
                "type": "object",
                "properties": {
                    "area": {"type": "string"},
                    "type": {"type": "string", "enum": ["need", "strength"]},
                    "severity": {"type": "string", "enum": ["high", "medium", "low"]},
                    "quote": {"type": "string"},
                    "explanation": {"type": "string"},
                    "suggested_support": {"type": "string", "enum": SUPPORT_OPTIONS},
                },
                "required": ["area", "type", "severity", "quote", "explanation", "suggested_support"],
            },
        },
    },
    "required": ["findings"],
}

SYSTEM = (
    "You read documents about one middle-school student for the school's student support office. "
    "Find the specific things this student needs help with, and the things they are good at.\n\n"
    "Rules:\n"
    "- Every finding needs `quote`: words copied EXACTLY from the document, 8 to 30 words, "
    "that show it. Copy, do not paraphrase. A finding you cannot quote does not exist.\n"
    "- `area` is the specific skill or behaviour, in a few plain words — 'multi-step word problems', "
    "not 'math'.\n"
    "- Describe what the student does. Never name a medical, psychological or learning condition, "
    "even if you suspect one.\n"
    "- Only findings about this student. Ignore anything about other students.\n"
    "- `suggested_support`: tutoring for a skill gap, homework-recovery when work is not being "
    "turned in, attendance-plan for absence, check-in for mood, effort or a broad slide, "
    "family-contact when home needs to know, enrichment for a strength, none otherwise.\n"
    "- If this part of the document says nothing useful, return no findings."
)


def _prompt(student_name: str, grade: int, kind: str, chunk: str, part: int, parts: int) -> str:
    where = f" This is part {part} of {parts}." if parts > 1 else ""
    return (f"Student: {student_name}, grade {grade}. Document type: {kind}.{where}\n\n"
            f"--- DOCUMENT ---\n{chunk}\n--- END ---")


@dataclass
class Outcome:
    findings: list[dict]
    rejected: list[dict]
    notices: list[str]
    summary: str


def review(raw_findings: list[dict], text: str, sig: StudentSignal | None) -> Outcome:
    """Apply the three checks. Pure: no model, no database — so it is fully testable."""
    norm = normalise(text)
    kept: list[dict] = []
    rejected: list[dict] = []
    notices: list[str] = []
    seen: set[tuple[str, str]] = set()

    for f in raw_findings:
        area = str(f.get("area") or "").strip()
        quote = str(f.get("quote") or "").strip()
        if not area or f.get("type") not in ("need", "strength"):
            rejected.append({**f, "reason": "Incomplete finding."})
            continue
        term = asserts_a_diagnosis(f)
        if term:
            rejected.append({**f, "reason": f"Names a condition ('{term}'). This tool records "
                                            "what a student does, not diagnoses."})
            note = ("The document may contain clinical or diagnostic language. Nothing diagnostic "
                    "was recorded — route that to the appropriate staff.")
            if note not in notices:
                notices.append(note)
            continue
        if not quote_in_document(quote, norm):
            rejected.append({**f, "reason": "The quoted words do not appear in the document."})
            continue
        key = (normalise(area), normalise(quote)[:60])
        if key in seen:
            continue
        seen.add(key)
        support = f.get("suggested_support") if f.get("suggested_support") in SUPPORT_OPTIONS else "none"
        if f["type"] == "strength" and support not in ("enrichment", "none"):
            support = "enrichment"
        item = {
            "area": area, "type": f["type"],
            "severity": f.get("severity") if f.get("severity") in SEVERITY_ORDER else "medium",
            "quote": quote, "explanation": str(f.get("explanation") or "").strip(),
            "suggested_support": support,
            "gradebook": corroborate(f, sig),
        }
        item["severity"], item["severity_note"] = calibrate_severity(item)
        kept.append(item)

    order = {"high": 0, "medium": 1, "low": 2}
    kept.sort(key=lambda f: (f["type"] != "need", order[f["severity"]]))
    return Outcome(findings=kept, rejected=rejected, notices=notices, summary=summarise(kept, rejected))


def summarise(findings: list[dict], rejected: list[dict]) -> str:
    """Written from the findings, not by the model — whose own summaries repeated the
    instructions back ("each finding includes an exact quote…") instead of the student."""
    needs = [f["area"] for f in findings if f["type"] == "need"]
    strengths = [f["area"] for f in findings if f["type"] == "strength"]
    if not findings:
        text = "Nothing specific enough to act on was found."
    else:
        parts = []
        if needs:
            parts.append("Needs help with " + ", ".join(needs[:4]) + ("…" if len(needs) > 4 else "") + ".")
        if strengths:
            parts.append("Strong at " + ", ".join(strengths[:3]) + ("…" if len(strengths) > 3 else "") + ".")
        text = " ".join(parts)
    if rejected:
        text += f" {len(rejected)} claim{'s' if len(rejected) != 1 else ''} withheld."
    return text


def analyse(db: Session, doc: StudentDocument) -> StudentDocument:
    started = time.time()
    student = doc.student
    sig = build_signals(db).get(student.sid)
    chunks = chunk_text(doc.text)
    if len(chunks) > MAX_CHUNKS:
        raise DocumentError(
            f"This document is about {len(doc.text):,} characters. Upload the part about "
            f"{student.name} — anything past {MAX_CHUNKS * CHUNK_CHARS:,} characters is too long "
            "to read reliably on the local model.")

    raw: list[dict] = []
    notices: list[str] = []
    for i, chunk in enumerate(chunks, start=1):
        reply = chat(
            [{"role": "system", "content": SYSTEM},
             {"role": "user", "content": _prompt(student.name, student.grade, doc.kind, chunk, i, len(chunks))}],
            fmt=SCHEMA,
        )
        if reply.truncated:
            notices.append(f"Part {i} may have been cut short by the model's context window, "
                           "so findings from it could be incomplete.")
        try:
            payload = json.loads(reply.content or "{}")
        except json.JSONDecodeError:
            notices.append(f"Part {i} of the document could not be read by the model.")
            continue
        for f in (payload.get("findings") or [])[:MAX_FINDINGS_PER_CHUNK]:
            if isinstance(f, dict):
                raw.append(f)

    outcome = review(raw, doc.text, sig)
    doc.summary = outcome.summary
    doc.findings = outcome.findings
    doc.rejected = outcome.rejected
    doc.notices = notices + outcome.notices
    doc.chunks = len(chunks)
    doc.model = settings.ollama_model
    doc.status = "done"
    doc.error = None
    doc.duration_ms = int((time.time() - started) * 1000)
    doc.analysed_at = datetime.utcnow()
    db.commit()
    return doc


def analyse_in_background(doc_id: int) -> None:
    from ..db import SessionLocal

    db = SessionLocal()
    try:
        doc = db.get(StudentDocument, doc_id)
        if doc is None:
            return
        try:
            analyse(db, doc)
        except (OllamaError, DocumentError) as e:
            db.rollback()
            doc = db.get(StudentDocument, doc_id)
            doc.status, doc.error = "failed", str(e)
            doc.analysed_at = datetime.utcnow()
            db.commit()
    except Exception as e:                        # never leave a document stuck on "processing"
        db.rollback()
        doc = db.get(StudentDocument, doc_id)
        if doc is not None:
            doc.status, doc.error = "failed", f"{type(e).__name__}: {e}"
            doc.analysed_at = datetime.utcnow()
            db.commit()
    finally:
        db.close()
