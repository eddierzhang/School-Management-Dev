#!/usr/bin/env python
"""Generate sample documents for trying the document reader.

Everything here is FICTIONAL — Halverson Ridge and its students are invented.
The note about Talia Barnard is written to match her seeded gradebook, and it
deliberately includes things the reader must handle carefully: a parent asking
about a condition by name (which must not be recorded as a finding) and a
mention of another student (whose details must not be attributed to Talia).

    python samples/make_samples.py     # writes .txt, .docx and .pdf next to this file
"""
from __future__ import annotations

import textwrap
from pathlib import Path

HERE = Path(__file__).parent

TITLE = "Teacher note — Talia Barnard (S-1507), Grade 7"
PARAGRAPHS = [
    "FICTIONAL SAMPLE DOCUMENT. Halverson Ridge Middle School, Fall 2026, week 5. "
    "Written by S. Frankel (Algebra Foundations) with input from L. Moreau (Digital Photography).",

    "Algebra. Talia can usually carry out each step of a problem on its own, but she loses track "
    "of what the question is asking partway through multi-step word problems. On the last quiz she "
    "set up two equations correctly and then answered a different question from the one on the page. "
    "She has not turned in three of the last five algebra assignments, and when I asked, she said "
    "she did not know where to start once she got home.",

    "Attendance. Her absences are clustered on Mondays, and she has missed several of our lunchtime "
    "review sessions. When she is here she participates, so the gap is time in the room rather than "
    "effort.",

    "At conferences her mother asked whether Talia might have ADHD, since she finds it hard to settle "
    "into written work at home. I said that was a question for the counselor and the family doctor, "
    "not for me.",

    "Photography. Ms. Moreau reports that her photographs show a real eye for composition and exposure, "
    "and she is often the first to help classmates with their camera settings. Unlike Marcus, who rushes "
    "his edits, Talia takes her time and asks good questions.",

    "Writing. Her short stories in Creative Writing have a clear narrative structure with a strong "
    "opening, even when her spelling and punctuation slip.",

    "Suggested next step: a short weekly check-in, and some help breaking word problems into parts.",
]


def write_txt() -> Path:
    p = HERE / "talia-barnard-teacher-note.txt"
    p.write_text(TITLE + "\n\n" + "\n\n".join(PARAGRAPHS) + "\n", encoding="utf-8")
    return p


def write_docx() -> Path:
    import docx

    d = docx.Document()
    d.add_heading(TITLE, level=1)
    for para in PARAGRAPHS:
        d.add_paragraph(para)
    p = HERE / "talia-barnard-teacher-note.docx"
    d.save(p)
    return p


def make_pdf(lines: list[str]) -> bytes:
    """A minimal one-or-more-page PDF with real, extractable text — no extra dependency."""
    def esc(s: str) -> str:
        return (s.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
                 .encode("latin-1", "replace").decode("latin-1"))

    per_page = 46
    pages = [lines[i:i + per_page] for i in range(0, len(lines), per_page)] or [[]]
    objects: list[bytes] = []
    font_id = 3 + 2 * len(pages)
    kids = " ".join(f"{3 + 2 * i} 0 R" for i in range(len(pages)))
    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objects.append(f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>".encode())
    for i, page_lines in enumerate(pages):
        content_id = 4 + 2 * i
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 {font_id} 0 R >> >> /Contents {content_id} 0 R >>".encode())
        ops = ["BT", "/F1 11 Tf", "14 TL", "56 736 Td"]
        for line in page_lines:
            ops.append(f"({esc(line)}) Tj T*")
        ops.append("ET")
        stream = "\n".join(ops).encode("latin-1", "replace")
        objects.append(b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream")
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")

    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for n, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{n} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return bytes(out)


def write_pdf() -> Path:
    lines = [TITLE.replace("—", "-"), ""]
    for para in PARAGRAPHS:
        lines += textwrap.wrap(para, 92) + [""]
    p = HERE / "talia-barnard-teacher-note.pdf"
    p.write_bytes(make_pdf(lines))
    return p


if __name__ == "__main__":
    for path in (write_txt(), write_docx(), write_pdf()):
        print(f"wrote {path.name} ({path.stat().st_size:,} bytes)")
